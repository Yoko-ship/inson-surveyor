import io
import json
import sqlite3
from datetime import date
from unittest.mock import Mock

import pytest
from PIL import Image
from sqlalchemy import select

from surveyor.calculations import calculate, valuation
from surveyor.db import Audit, Channel, Indicator, now
from surveyor.maintenance import create_backup, restore_backup, verify_backup
from surveyor.schemas import IndicatorInput
from surveyor.source_adapters import AccessRefused, checked_url, parse_public
from surveyor.sources import collect_channel, latest_indicators, store_indicator
from tests.conftest import switch_user


def item(**kwargs):
    return IndicatorInput(
        metric="crime",
        period="2025",
        value="50",
        unit="count",
        source_url="https://stat.uz/",
        observation_date=date(2025, 12, 31),
        **kwargs,
    )


def test_old_period_does_not_replace_latest_and_approval_persists(admin):
    with admin.factory() as db:
        row = store_indicator(db, "stat", item(region="1726"))
        db.commit()
        row.data = {**row.data, "approved_by": "actuary", "approved_at": now().isoformat()}
        db.commit()
        again = store_indicator(db, "stat", item(region="1726"))
        assert again.id == row.id
        data = item(region="1726").model_dump(mode="json")
        data.update(period="2024", observation_date="2024-12-31", value="20")
        store_indicator(db, "stat", data)
        db.commit()
        rows = latest_indicators(db, "Ташкент")
        assert len(rows) == 1
        assert rows[0]["value"] == "50"
        assert rows[0]["approved_by"] == "actuary"


@pytest.mark.parametrize(
    "url",
    [
        "http://stat.uz/a",
        "https://evil.test/a",
        "https://stat.uz.evil.test/a",
        "https://user:secret@stat.uz/a",
        "https://127.0.0.1/a",
        "https://stat.uz:444/a",
    ],
)
def test_source_url_boundary(url):
    with pytest.raises(ValueError):
        checked_url(url, "stat", resolve=False)


def test_contract_only_sources_cannot_enable_collection(admin):
    config = {
        "url": "https://stat.uz/public.csv",
        "format": "csv",
        "permission_url": "https://stat.uz/",
        "permission_note": "Official public download terms",
        "columns": {"value": "value"},
        "constants": {"metric": "x", "period": "2025", "unit": "count", "observation_date": "2025-12-31"},
    }
    assert admin.put("/api/admin/sources/credit/config", json=config).status_code == 422
    assert (
        admin.patch(
            "/api/admin/sources/listings", json={"enabled": True, "reason": "Reviewed public access"}
        ).status_code
        == 422
    )
    assert admin.put("/api/admin/sources/stat/config", json=config).status_code == 200
    switch_user(admin, "employee")
    assert admin.put("/api/admin/sources/stat/config", json=config).status_code == 403


def test_csv_and_html_mapping_and_schema():
    config = {
        "format": "csv",
        "url": "https://stat.uz/file.csv",
        "columns": {"value": "Count", "region": "Region"},
        "constants": {"metric": "crime", "period": "2025", "unit": "count", "observation_date": "2025-12-31"},
    }
    rows, schema = parse_public(b"Count,Region\n25,1726\n", config)
    assert rows[0].value == 25 and schema == ["Count", "Region"]
    config["format"] = "html_table"
    rows, _ = parse_public(
        b"<table><tr><th>Count</th><th>Region</th></tr><tr><td>25</td><td>1726</td></tr></table>", config
    )
    assert rows[0].value == 25
    with pytest.raises(ValueError):
        parse_public(b'<table><tr><td colspan="2">Changed</td></tr></table>', config)


def test_siat_periods_provenance_and_future_guard():
    config = {"format": "siat", "url": "https://api.siat.stat.uz/file.json", "constants": {"metric": "crime"}}
    payload = [
        {
            "metadata": [{"name_en": "Unit of measurement", "value_en": "count"}],
            "data": [{"Code": "1726", "Klassifikator_ru": "Ташкент", "2024": 5, "2025": 10, "2099": 999}],
        }
    ]
    rows, _ = parse_public(json.dumps(payload).encode(), config)
    assert [r.period for r in rows] == ["2024", "2025"]
    assert rows[-1].observation_date == date(2025, 12, 31)
    assert rows[-1].region == "1726" and rows[-1].rate_adjustment == 0


def test_collector_disables_on_refusal_and_never_retries(admin, monkeypatch):
    from surveyor import source_adapters

    fetch = Mock(side_effect=AccessRefused("Access refused"))
    monkeypatch.setattr(source_adapters, "fetch_public", fetch)
    with admin.factory() as db:
        assert collect_channel(db, "stat_crime")["status"] == "error"
        assert not db.get(Channel, "stat_crime").enabled
        assert collect_channel(db, "stat_crime", force=True)["status"] == "disabled"
        assert fetch.call_count == 1
        assert db.scalar(select(Audit).where(Audit.action == "source.error"))


def test_collector_schema_change_does_not_partially_write(admin, monkeypatch):
    from surveyor import source_adapters

    with admin.factory() as db:
        c = db.get(Channel, "stat_crime")
        config = {
            "format": "csv",
            "url": "https://stat.uz/file.csv",
            "columns": {"value": "Count"},
            "constants": {
                "metric": "crime",
                "period": "2025",
                "unit": "count",
                "observation_date": "2025-12-31",
            },
        }
        c.data = {**c.data, "config": config, "schema": ["Count", "Old"]}
        db.commit()
        monkeypatch.setattr(source_adapters, "fetch_public", lambda *args: b"Count,Changed\n25,a\n")
        assert collect_channel(db, "stat_crime")["status"] == "error"
        assert not db.scalars(select(Indicator)).all()
        assert not c.enabled


def test_indicator_import_preview_confirm_atomic_and_owned(admin):
    csv = b"metric,period,value,unit,source_url,observation_date\ncrime,2025,15,count,https://stat.uz/,2025-12-31\n"
    batch = admin.post(
        "/api/admin/sources/stat/imports/preview", files={"file": ("indicators.csv", csv)}
    ).json()
    assert batch["can_confirm"]
    assert admin.get("/api/sources").json()["indicators"] == []
    assert admin.post(f"/api/admin/source-imports/{batch['id']}/confirm").status_code == 200
    assert admin.post(f"/api/admin/source-imports/{batch['id']}/confirm").status_code == 409
    assert len(admin.get("/api/sources").json()["indicators"]) == 1
    invalid = admin.post(
        "/api/admin/sources/stat/imports/preview",
        files={"file": ("bad.csv", csv + b"crime,2025,-2,count,https://stat.uz/,2025-12-31\n")},
    ).json()
    assert not invalid["can_confirm"]
    assert admin.post(f"/api/admin/source-imports/{invalid['id']}/confirm").status_code == 422
    assert len(admin.get("/api/sources").json()["indicators"]) == 1


def test_document_manual_review_retains_original_and_invalidates_confirmation(admin):
    survey = admin.post("/api/surveys", json={"title": "Manual review"}).json()
    doc = admin.post(
        f"/api/surveys/{survey['id']}/documents",
        files={"file": ("contract.txt", "Договор\nСтраховая сумма: 1000".encode())},
    ).json()
    body = {
        "revision": doc["revision"],
        "kind": "contract",
        "fields": {"insured_sum": "1200"},
        "reason": "Corrected after visual review",
    }
    r = admin.put(f"/api/documents/{doc['id']}/review", json=body)
    assert r.status_code == 200, r.text
    field = r.json()["extracted"]["fields"]["insured_sum"]
    assert field["original"] == "1000" and field["value"] == "1200" and field["status"] == "manual"
    assert admin.put(f"/api/documents/{doc['id']}/review", json=body).status_code == 409
    assert admin.post(f"/api/surveys/{survey['id']}/reports").status_code == 422
    switch_user(admin, "employee")
    assert (
        admin.put(
            f"/api/documents/{doc['id']}/review", json={**body, "revision": r.json()["revision"]}
        ).status_code
        == 404
    )


def test_scan_can_be_manually_structured_without_ai(admin):
    buffer = io.BytesIO()
    Image.new("RGB", (10, 10)).save(buffer, format="PNG")
    survey = admin.post("/api/surveys", json={"title": "Scan"}).json()
    doc = admin.post(
        f"/api/surveys/{survey['id']}/documents", files={"file": ("scan.png", buffer.getvalue())}
    ).json()
    r = admin.put(
        f"/api/documents/{doc['id']}/review",
        json={
            "revision": doc["revision"],
            "kind": "branch_request",
            "fields": {"declared_rate": "0.5", "insured_organization": "LLC Example"},
            "reason": "Read from scan by employee",
        },
    )
    assert r.status_code == 200
    assert r.json()["extracted"]["original_kind"] == "photo_or_scan"


def test_valuation_threshold_requires_second_method_and_configurable_outliers():
    inputs = {
        "object_type": "vehicle",
        "object_value": "1000",
        "comparables": [{"label": "A", "source": "file", "price": "1000", "date": date.today().isoformat()}],
    }
    v = valuation(inputs, policy={"large_object_threshold": "500"})
    assert v["estimate"] is None and v["policy"]["requires_appraiser"]
    v = valuation({**inputs, "object_value": "1200"}, policy={"valuation_tolerance": "0.25"})
    assert v["status"] == "confirmed"


def test_regional_rule_bounded_object_scoped_and_floor():
    product = {"rate_type": "annual", "rate": "0.5", "min_rate": "0.4", "class_code": "property"}
    inputs = {"insured_sum": "100000", "region": "1726", "object_type": "housing"}
    template = {
        "moderate_threshold": 20,
        "high_threshold": 50,
        "multipliers": {"low": "1", "moderate": "1", "high": "1"},
        "max_adjustment": "0.1",
        "indicator_rules": {
            "crime": {
                "baseline": "100",
                "sensitivity": "1",
                "max_adjustment": "0.4",
                "object_types": ["housing"],
            }
        },
    }
    indicators = [{"metric": "crime", "value": "500", "object_type": "housing"}]
    c = calculate(product, inputs, template, indicators)
    assert c["regional_adjustment"] == "0.1" and c["recommended_rate"] == "0.55"
    assert (
        calculate(product, {**inputs, "object_type": "equipment"}, template, indicators)[
            "regional_adjustment"
        ]
        == "0"
    )
    assert (
        calculate(
            {
                **product,
                "rate_type": "normative",
                "normative_basis": "annual",
                "normative_source": "https://lex.uz/",
            },
            inputs,
            template,
            indicators,
        )["recommended_rate"]
        == "0.5"
    )


def test_backup_restore_checks_documents_and_rejects_overwrite(tmp_path):
    import hashlib

    from sqlalchemy import create_engine

    from surveyor.db import Base

    source = tmp_path / "live.sqlite"
    engine = create_engine(f"sqlite:///{source}")
    Base.metadata.create_all(engine)
    engine.dispose()
    doc = tmp_path / "source.txt"
    doc.write_text("source evidence", encoding="utf-8")
    with sqlite3.connect(source) as conn:
        conn.execute(
            "INSERT INTO users (id,login,phone,name,position,department,branch,role,password_hash,must_change_password,active,created_at) VALUES ('u','u','1','A','','','','admin','hash',0,1,'2026-01-01')"
        )
        conn.execute(
            "INSERT INTO surveys (id,owner_id,title,status,inputs,revision,created_at) VALUES ('s','u','Test','draft','{}',1,'2026-01-01')"
        )
        conn.execute(
            "INSERT INTO documents (id,survey_id,filename,path,sha256,extracted,created_at) VALUES ('d','s','source.txt',?,?,'{}','2026-01-01')",
            (str(doc), hashlib.sha256(doc.read_bytes()).hexdigest()),
        )
        conn.commit()
    from surveyor import ai_config

    config = ai_config.defaults()
    backup = create_backup(f"sqlite:///{source}", tmp_path / "backups")
    assert verify_backup(backup)["documents"] == {"d": "uploads/d.txt"}
    restored = restore_backup(backup, tmp_path / "restored")
    assert (
        ai_config.AIConfig.model_validate_json((restored / "ai/config.json").read_text(encoding="utf-8"))
        == config
    )
    assert (restored / "ai/guardrails.txt").read_text(encoding="utf-8") == ai_config.BASELINE_PATH.read_text(
        encoding="utf-8"
    )
    with sqlite3.connect(restored / "database.sqlite") as conn:
        from pathlib import Path

        path = conn.execute("SELECT path FROM documents").fetchone()[0]
        assert Path(path).read_text(encoding="utf-8") == "source evidence"
    with pytest.raises(ValueError):
        restore_backup(backup, restored)
    (backup / "uploads/d.txt").write_text("tampered", encoding="utf-8")
    with pytest.raises(ValueError):
        verify_backup(backup)


@pytest.mark.parametrize("language,title", [("uz", "Obyekt va hujjatlar"), ("en", "Object and documents")])
def test_localized_screen_and_exports_share_five_sections(admin, language, title):
    survey = admin.post("/api/surveys", json={"title": "Example"}).json()
    inputs = {
        "revision": 1,
        "product_code": "DEMO-FIX",
        "insured_sum": "1000000",
        "object_value": "1000000",
        "region": "1726",
        "language": language,
        "manual_review_confirmed": True,
    }
    assert admin.put(f"/api/surveys/{survey['id']}", json=inputs).status_code == 200
    report = admin.post(f"/api/surveys/{survey['id']}/reports").json()
    detail = admin.get(f"/api/reports/{report['id']}").json()
    assert len(detail["sections"]) == 5 and title in detail["sections"][0]["title"]
    assert "Стоимость объекта" not in "\n".join(detail["sections"][0]["lines"])
    assert admin.get(f"/api/reports/{report['id']}/export/docx").status_code == 200
    assert admin.get(f"/api/reports/{report['id']}/export/pdf").status_code == 200


def test_public_reference_versions_and_report_snapshot(admin):
    body = {
        "title": "Example regulation",
        "source_url": "https://lex.uz/docs/123",
        "text": "Original regulation text",
        "kind": "law",
        "observation_date": "2025-12-31",
    }
    first = admin.post("/api/admin/sources/lex/references", json=body)
    assert first.status_code == 201, first.text
    old_id = first.json()["id"]
    assert admin.post("/api/admin/sources/lex/references", json=body).json()["id"] == old_id
    survey = admin.post("/api/surveys", json={"title": "Referenced report"}).json()
    admin.put(
        f"/api/surveys/{survey['id']}",
        json={
            "revision": 1,
            "product_code": "DEMO-FIX",
            "insured_sum": "100",
            "object_value": "100",
            "region": "1726",
            "reference_ids": [old_id],
            "manual_review_confirmed": True,
        },
    )
    report = admin.post(f"/api/surveys/{survey['id']}/reports").json()
    new = admin.post("/api/admin/sources/lex/references", json={**body, "text": "Amended regulation text"})
    assert new.status_code == 201 and new.json()["id"] != old_id
    assert len(admin.get("/api/admin/sources/lex/references/history").json()) == 2
    snapshot = admin.get(f"/api/reports/{report['id']}").json()["snapshot"]
    assert snapshot["references"][0]["text"] == "Original regulation text"
    assert admin.get("/api/sources").json()["references"][0]["text"] == "Amended regulation text"
    assert any(
        a["message"] == "Example regulation" for a in admin.get("/api/admin/operations").json()["alerts"]
    )


def test_document_monitor_does_not_treat_text_as_rate(admin, monkeypatch):
    from surveyor import source_adapters

    config = {
        "url": "https://lex.uz/docs/123",
        "format": "document",
        "permission_url": "https://lex.uz/",
        "permission_note": "Document automation approved for testing",
        "reference": {"title": "Act", "kind": "law"},
    }
    assert admin.put("/api/admin/sources/lex/config", json=config).status_code == 200
    monkeypatch.setattr(
        source_adapters,
        "fetch_public",
        lambda *args: b"<html><script>alert(1)</script><p>Regulation: 50%</p></html>",
    )
    result = admin.post("/api/admin/sources/lex/collect")
    assert result.json()["status"] == "collected"
    data = admin.get("/api/sources").json()
    assert data["indicators"] == []
    assert data["references"][0]["text"] == "Regulation: 50%"
    assert data["references"][0]["stale"]


@pytest.mark.parametrize("period,expected", [("2025-Q1", "2025-03-31"), ("2025-M02", "2025-02-28")])
def test_siat_quarter_and_month_period_end(period, expected):
    payload = [
        {
            "metadata": [{"name_en": "Unit of measurement", "value_en": "count"}],
            "data": [{"Code": "1726", "Klassifikator_ru": "Ташкент", period: 10}],
        }
    ]
    rows, _ = parse_public(
        json.dumps(payload).encode(),
        {"format": "siat", "url": "https://stat.uz/data.json", "constants": {"metric": "test"}},
    )
    assert rows[0].observation_date.isoformat() == expected


def napp_workbook():
    from datetime import datetime

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "1.4"
    ws.cell(3, 2, "Umumiy sug‘urta mukofotlari")
    ws.cell(3, 5, "Sug‘urta to‘lovlari")
    ws.cell(3, 8, "Sug‘urta majburiyatlari")
    for col in [2, 5, 8]:
        ws.cell(4, col, datetime(2025, 7, 1))
        ws.cell(4, col + 1, datetime(2026, 7, 1))
        ws.cell(5, col, "mln. so‘mda")
    for cls in range(1, 19):
        row = cls + 5
        ws.cell(row, 1, f"{cls}-klass Example")
        for col, value in [(2, 10), (3, 20), (5, 2), (6, 5), (8, 1000), (9, 2000)]:
            ws.cell(row, col, value)
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


def test_napp_official_periods_units_and_no_invented_market_rate():
    from surveyor.napp import latest_download, parse_napp

    url = latest_download(
        b'<a href="/storage/files/shares/opendata/2026/2Q/report.xlsx">XLSX</a><a href="/storage/files/shares/opendata/2099/3Q/future.xlsx">Future</a>'
    )
    assert "/2026/2Q/" in url
    items, _ = parse_napp(napp_workbook(), url)
    vehicle = [i for i in items if i.class_code == "vehicle" and i.period.startswith("2026")]
    assert {i.metric: str(i.value) for i in vehicle} == {
        "market_premiums": "20000000",
        "market_payments": "5000000",
        "market_liabilities": "2000000000",
        "market_loss_ratio": "0.25",
    }
    assert all(i.annual_market_rate is None for i in items)
    assert all(i.observation_date == date(2026, 6, 30) for i in vehicle)


def test_robots_duplicate_groups_block_disallowed_path(monkeypatch, tmp_path):
    import httpx

    from surveyor import source_adapters
    from surveyor.config import settings

    calls = []

    def handler(request):
        calls.append(str(request.url))
        return httpx.Response(200, text="User-agent: *\nAllow: /\n\nUser-agent: *\nDisallow: /page\n")

    real_client = httpx.Client
    monkeypatch.setattr(
        source_adapters.httpx,
        "Client",
        lambda **kw: real_client(transport=httpx.MockTransport(handler), **kw),
    )
    monkeypatch.setattr(source_adapters, "checked_url", lambda url, channel: url)
    monkeypatch.setattr(settings, "storage_dir", tmp_path / "uploads")
    with pytest.raises(AccessRefused):
        source_adapters.fetch_public({"url": "https://uzex.uz/pages/quotes"}, "exchange")
    assert calls == ["https://uzex.uz/robots.txt"]


def test_corrupt_external_workbook_disables_channel(admin, monkeypatch):
    from surveyor import source_adapters

    responses = iter(
        [b'<a href="/storage/files/shares/opendata/2026/2Q/report.xlsx">XLSX</a>', b"not a workbook"]
    )
    monkeypatch.setattr(source_adapters, "fetch_public", lambda *args: next(responses))
    with admin.factory() as db:
        result = collect_channel(db, "napp_market")
        assert result["status"] == "error"
        assert not db.get(Channel, "napp_market").enabled


def test_empty_excel_and_duplicate_csv_headers_are_clear_errors(admin):
    from openpyxl import Workbook

    stream = io.BytesIO()
    Workbook().save(stream)
    for filename, data in [("empty.xlsx", stream.getvalue()), ("duplicate.csv", b"metric,metric\na,b\n")]:
        response = admin.post("/api/admin/sources/stat/imports/preview", files={"file": (filename, data)})
        assert response.status_code == 422, response.text
        assert "detail" in response.json()
