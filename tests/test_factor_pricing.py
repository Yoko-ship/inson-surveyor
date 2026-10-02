import copy
import io
from datetime import date
from decimal import Decimal

import pytest
from conftest import switch_user
from docx import Document
from pydantic import ValidationError

from surveyor.calculations import calculate
from surveyor.factor_pricing import catalogue, entries, propose
from surveyor.schemas import FactorPolicy, SurveyInput


def policy():
    return {
        "insurance_class": 3,
        "coefficients": {
            "factor_001": {"raises": "1.2", "lowers": "0.8"},
            "factor_002": {"raises": "1.5", "lowers": "0.9"},
        },
        "rationale": "Fictional coefficients for automated acceptance",
    }


def template(**kw):
    return {
        "id": "version-one",
        "feature_weights": {"damage": 80},
        "moderate_threshold": 20,
        "high_threshold": 50,
        "multipliers": {"low": "1", "moderate": "2", "high": "3"},
        "approved_by": "actuary",
        "factor_policy": policy(),
        **kw,
    }


def product(**kw):
    return {
        "code": "P",
        "class_code": "vehicle",
        "rate": "0.5",
        "min_rate": "0.3",
        "rate_type": "annual",
        **kw,
    }


def answers(complete=False):
    result = (
        {r["id"]: {"choice": "not_applicable", "evidence": "Fictional scope reviewed"} for r in entries(3)}
        if complete
        else {}
    )
    result.update(
        {
            "factor_001": {"choice": "raises", "evidence": "Fictional loss record"},
            "factor_002": {"choice": "lowers", "evidence": "Fictional complete documents"},
        }
    )
    return result


def inputs(**kw):
    return {
        "insured_sum": "100000000",
        "object_value": "100000000",
        "term_days": 365,
        "factor_answers": answers(),
        "features": ["damage"],
        **kw,
    }


@pytest.mark.parametrize(
    "kind,basis,expected",
    [
        ("annual", None, "1620000.00"),
        ("fixed", None, "540000.00"),
        ("program", "annual", "1944000.00"),
        ("program", "fixed", "648000.00"),
        ("normative", "fixed", "500000.00"),
        ("normative", "annual", "1500000.00"),
    ],
)
def test_exact_factor_pricing_all_bases_without_double_count(kind, basis, expected):
    p = product(
        rate_type=kind,
        program_rates={"A": "0.6"},
        program_basis=basis,
        normative_basis=basis,
        normative_source="https://example.test/norm",
    )
    calc = calculate(
        p,
        inputs(term_days=1095, program="A"),
        template(),
        [{"metric": "regional_risk", "rate_adjustment": "0.2"}],
        {"adjustment": "0.2"},
    )
    assert calc["premium"] == expected
    assert calc["regional_adjustment"] == calc["loss_adjustment"] == "0"
    assert calc["factor_pricing"]["multiplier"] == ("1" if kind == "normative" else "1.08")


def test_minimum_missing_and_null_coefficients():
    t = template()
    t["factor_policy"]["coefficients"]["factor_001"]["lowers"] = "0.1"
    a = {
        "factor_001": {"choice": "lowers", "evidence": "Fictional evidence"},
        "factor_003": {"choice": "raises", "evidence": "Fictional unpriced factor"},
    }
    c = calculate(product(), inputs(factor_answers=a), t)
    assert c["recommended_rate"] == "0.3"
    reasons = {r["id"]: r["reason"] for r in c["factor_pricing"]["clarify"]}
    assert reasons["factor_003"] == "coefficient_missing"
    assert reasons["factor_002"] == "unanswered"
    assert not c["factor_pricing"]["ready_for_underwriting"]


@pytest.mark.parametrize("changes", [{"approved_by": None}, {"factor_calibration_stale": True}])
def test_unapproved_or_stale_not_applied(changes):
    c = calculate(product(), inputs(), template(**changes))
    assert c["premium"] == "500000.00"
    assert c["factor_pricing"]["proposed_multiplier"] == "1.08"
    assert not c["factor_pricing"]["ready_for_underwriting"]


@pytest.mark.parametrize(
    "change",
    [
        {"insurance_class": 19},
        {"coefficients": {"unknown": {"raises": "1.2"}}},
        {"coefficients": {"factor_001": {"raises": "0.9"}}},
        {"coefficients": {"factor_001": {"lowers": "1"}}},
        {"coefficients": {"factor_001": {"raises": "NaN"}}},
        {"coefficients": {"factor_001": {"raises": "Infinity"}}},
        {"coefficients": {"factor_001": {"raises": "10.1"}}},
    ],
)
def test_coefficient_schema_rejects_invalid(change):
    with pytest.raises(ValidationError):
        FactorPolicy.model_validate({**policy(), **change})


def test_catalogue_complete_and_foreign_factor_rejected():
    assert len(catalogue()["entries"]) == 179
    for cls in range(1, 19):
        assert entries(cls)
    assert len([r for r in entries(4) if r["id"].startswith("note_")]) == 4
    other = next(r for r in entries(8) if r not in entries(3))
    with pytest.raises(ValueError, match="другого класса"):
        calculate(
            product(),
            inputs(factor_answers={other["id"]: {"choice": "raises", "evidence": "test"}}),
            template(),
        )
    with pytest.raises(ValueError, match="шаблон"):
        from surveyor.factor_pricing import validate_answers

        validate_answers(inputs(), None)


def create_policy(client):
    old = next(t for t in client.get("/api/templates").json() if t["class_code"] == "vehicle")
    response = client.post(f"/api/admin/templates/{old['id']}/factors", json={"policy": policy()})
    assert response.status_code == 201, response.text
    return response.json()["id"]


def experience_rows():
    return [
        {
            "factor_id": "factor_001",
            "choice": choice,
            "year": year,
            "exposure": "100",
            "claims": 10,
            "payments": payments,
        }
        for year in range(date.today().year - 3, date.today().year)
        for choice, payments in [("neutral", "1000"), ("raises", "1500"), ("lowers", "500")]
    ]


def upload(client, template_id, rows=None):
    rows = experience_rows() if rows is None else rows
    text = "factor_id,choice,year,exposure,claims,payments\n" + "\n".join(
        ",".join(str(r[k]) for k in ["factor_id", "choice", "year", "exposure", "claims", "payments"])
        for r in rows
    )
    response = client.post(
        f"/api/admin/templates/{template_id}/factor-experience/preview",
        files={"file": ("experience.csv", text.encode(), "text/csv")},
    )
    assert response.status_code == 200, response.text
    return response.json()


def test_version_approval_snapshot_exports_and_role_access(admin):
    tid = create_policy(admin)
    body = {
        "revision": 1,
        "product_code": "DEMO-AUTO",
        "region": "1726",
        "tariff_date": str(date.today()),
        "insured_sum": "100000000",
        "object_value": "100000000",
        "manual_review_confirmed": True,
        "factor_answers": answers(complete=True),
    }
    assert admin.post("/api/calculate", json=body).json()["premium"] == "500000.00"
    assert admin.post(f"/api/admin/templates/{tid}/approve").status_code == 403
    switch_user(admin, "actuary")
    assert admin.post(f"/api/admin/templates/{tid}/approve").status_code == 200
    assert admin.post("/api/calculate", json=body).json()["premium"] == "540000.00"
    sid = admin.post("/api/surveys", json={"title": "FICTIONAL factor test"}).json()["id"]
    assert admin.put(f"/api/surveys/{sid}", json=body).status_code == 200
    report = admin.post(f"/api/surveys/{sid}/reports").json()
    rid = report["id"]
    snapshot = admin.get(f"/api/reports/{rid}").json()["snapshot"]
    assert snapshot["calculation"]["factor_pricing"]["ready_for_underwriting"]
    for fmt in ["pdf", "docx"]:
        response = admin.get(f"/api/reports/{rid}/export/{fmt}")
        assert response.status_code == 200
        if fmt == "docx":
            text = "\n".join(p.text for p in Document(io.BytesIO(response.content)).paragraphs)
            assert "История убытков клиента" in text and "1.08" in text
    newer = admin.post(f"/api/admin/templates/{tid}/factors", json={"policy": policy()})
    assert newer.status_code == 201
    assert admin.get(f"/api/reports/{rid}").json()["snapshot"] == snapshot
    assert admin.post("/api/calculate", json=body).json()["premium"] == "500000.00"
    switch_user(admin, "underwriter", "underwriter")
    response = admin.post(
        f"/api/reports/{rid}/decision", json={"decision": "approved", "comment": "Fictional review reason"}
    )
    assert response.status_code == 409
    switch_user(admin, "employee", "employee-two")
    assert admin.post(f"/api/admin/templates/{tid}/factors", json={"policy": policy()}).status_code == 403
    assert admin.get(f"/api/admin/templates/{tid}/factor-experience").status_code == 403


def test_experience_proposal_approval_and_changed_data_invalidation(admin):
    tid = create_policy(admin)
    preview = upload(admin, tid)
    bid = preview["id"]
    assert preview["can_confirm"]
    assert admin.post(f"/api/admin/factor-experience/{bid}/confirm").status_code == 200
    assert admin.post(f"/api/admin/factor-experience/{bid}/confirm").status_code == 409
    request = {
        "batch_id": bid,
        "minimum_exposure": "100",
        "max_change": "0.2",
        "rationale": "Fictional actuarial method acceptance",
        "method_confirmed": True,
    }
    assert admin.post(f"/api/admin/templates/{tid}/factor-calibration", json=request).status_code == 403
    switch_user(admin, "actuary")
    response = admin.post(f"/api/admin/templates/{tid}/factor-calibration", json=request)
    assert response.status_code == 201, response.text
    result = response.json()
    assert result["policy"]["coefficients"]["factor_001"] == {"raises": "1.20000000", "lowers": "0.80000000"}
    new_id = result["id"]
    assert admin.post(f"/api/admin/templates/{new_id}/approve").status_code == 200
    rows = experience_rows()
    rows[0]["payments"] = "1100"
    changed = upload(admin, new_id, rows)
    assert admin.post(f"/api/admin/factor-experience/{changed['id']}/confirm").status_code == 200
    latest = next(t for t in admin.get("/api/templates").json() if t["class_code"] == "vehicle")
    assert latest["factor_calibration_stale"]
    assert admin.post(f"/api/admin/templates/{new_id}/approve").status_code == 409
    assert admin.post(f"/api/admin/templates/{new_id}/factor-calibration", json=request).status_code == 409


@pytest.mark.parametrize(
    "mutate", ["gap", "duplicate", "zero_neutral", "too_little", "direction", "wrong_class"]
)
def test_statistical_recalculation_rejects_bad_or_inadequate_data(mutate):
    rows = experience_rows()
    minimum = Decimal("100")
    if mutate == "gap":
        rows.pop()
    if mutate == "duplicate":
        rows.append(copy.deepcopy(rows[0]))
    if mutate == "zero_neutral":
        for r in rows:
            if r["choice"] == "neutral":
                r["payments"] = "0"
    if mutate == "too_little":
        minimum = Decimal("301")
    if mutate == "direction":
        for r in rows:
            if r["choice"] == "raises":
                r["payments"] = "500"
    if mutate == "wrong_class":
        for r in rows:
            r["factor_id"] = "note_4_1"
    with pytest.raises(ValueError):
        propose(policy(), rows, minimum, Decimal("0.2"))


def test_import_validation_atomic_csrf_and_template_staleness(admin):
    tid = create_policy(admin)
    bad = experience_rows()
    bad.append(bad[0])
    preview = upload(admin, tid, bad)
    assert not preview["can_confirm"]
    assert admin.post(f"/api/admin/factor-experience/{preview['id']}/confirm").status_code == 422
    good = upload(admin, tid)
    csrf = admin.headers.pop("X-CSRF-Token")
    assert admin.post(f"/api/admin/factor-experience/{good['id']}/confirm").status_code == 403
    admin.headers["X-CSRF-Token"] = csrf
    assert admin.post(f"/api/admin/templates/{tid}/factors", json={"policy": policy()}).status_code == 201
    assert admin.post(f"/api/admin/factor-experience/{good['id']}/confirm").status_code == 409


def test_answer_requires_evidence():
    with pytest.raises(ValidationError):
        SurveyInput.model_validate(
            {
                "revision": 1,
                "product_code": "P",
                "region": "1726",
                **inputs(),
                "factor_answers": {"factor_001": {"choice": "not_applicable", "evidence": ""}},
            }
        )


@pytest.mark.parametrize("years", [3, 4, 5])
def test_three_to_five_year_calibration_uses_exposure_not_claim_counts(years):
    rows = [
        {
            "factor_id": "factor_001",
            "choice": choice,
            "year": year,
            "exposure": str(exposure),
            "claims": claims,
            "payments": str(payments),
        }
        for year in range(date.today().year - years, date.today().year)
        for choice, exposure, claims, payments in [("neutral", 100, 1, 1000), ("raises", 200, 1, 3000)]
    ]
    result, estimates = propose(policy(), rows, Decimal(100), Decimal("0.9"))
    assert result["coefficients"]["factor_001"]["raises"] == "1.50000000"
    assert estimates[0]["raw"] == "1.5"


def test_factor_import_xlsx_and_foreign_owner(admin):
    from openpyxl import Workbook

    tid = create_policy(admin)
    book = Workbook()
    sheet = book.active
    keys = ["factor_id", "choice", "year", "exposure", "claims", "payments"]
    sheet.append(keys)
    for row in experience_rows():
        sheet.append([row[k] for k in keys])
    buffer = io.BytesIO()
    book.save(buffer)
    response = admin.post(
        f"/api/admin/templates/{tid}/factor-experience/preview",
        files={"file": ("experience.xlsx", buffer.getvalue())},
    )
    assert response.status_code == 200 and response.json()["can_confirm"]
    switch_user(admin, "actuary")
    assert admin.post(f"/api/admin/factor-experience/{response.json()['id']}/confirm").status_code == 404


def test_underwriter_can_approve_complete_current_factor_report(admin):
    tid = create_policy(admin)
    switch_user(admin, "actuary")
    assert admin.post(f"/api/admin/templates/{tid}/approve").status_code == 200
    sid = admin.post("/api/surveys", json={"title": "FICTIONAL approved factors"}).json()["id"]
    body = {
        "revision": 1,
        "product_code": "DEMO-AUTO",
        "region": "1726",
        "insured_sum": "100000000",
        "object_value": "100000000",
        "manual_review_confirmed": True,
        "factor_answers": answers(complete=True),
    }
    assert admin.put(f"/api/surveys/{sid}", json=body).status_code == 200
    rid = admin.post(f"/api/surveys/{sid}/reports").json()["id"]
    switch_user(admin, "underwriter", "factor-underwriter")
    response = admin.post(
        f"/api/reports/{rid}/decision", json={"decision": "approved", "comment": "Fictional acceptance"}
    )
    assert response.status_code == 200, response.text


def test_extreme_factor_product_rejected_instead_of_overflow():
    t = template()
    for key in ["factor_001", "factor_002", "factor_003", "factor_004"]:
        t["factor_policy"]["coefficients"][key] = {"raises": "10"}
    a = {
        k: {"choice": "raises", "evidence": "Fictional risk evidence"}
        for k in t["factor_policy"]["coefficients"]
    }
    with pytest.raises(ValueError, match="100%"):
        calculate(product(), inputs(factor_answers=a), t)


def test_prose_factors_retain_pdf_labels_and_pages_including_class_nine():
    source = catalogue()
    notes = [r for r in source["entries"] if r["id"].startswith("note_")]
    assert len(notes) == 40
    assert len([r for r in notes if r["classes"] == [9]]) == 6
    for row in notes:
        text = next(p["text"] for p in source["pages"] if p["page"] == row["page"])
        assert row["label"].lower() in " ".join(text.lower().split())


def test_numeric_template_class_cannot_select_another_legal_class(admin):
    payload = {
        "class_code": "3",
        "name": "Fictional class 3",
        "multipliers": {"low": "1", "moderate": "1", "high": "1"},
    }
    response = admin.post("/api/admin/templates", json=payload)
    assert response.status_code == 201
    tid = response.json()["id"]
    wrong = {**policy(), "insurance_class": 8}
    assert admin.post(f"/api/admin/templates/{tid}/factors", json={"policy": wrong}).status_code == 422
    assert admin.post("/api/admin/templates", json={**payload, "factor_policy": wrong}).status_code == 422
