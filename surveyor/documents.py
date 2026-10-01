import csv
import io
import re
import zipfile
from datetime import datetime
from pathlib import Path

from docx import Document as WordDocument
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader

from surveyor.document_values import derive_document_term, document_number

FIELDS = {
    "insured_sum": [
        r"страхов(?:ая|ую)\s+сумм[ауы]",
        r"sug['‘’]?urta\s+summasi",
        r"суғурта\s+суммаси",
        r"insured\s+sum",
    ],
    "object_value": [r"стоимость\s+объекта", r"объект\s+стоимость", r"obyekt\s+qiymati", r"object\s+value"],
    "declared_premium": [
        r"страхов(?:ая|ую)\s+преми[яю]",
        r"sug['‘’]?urta\s+mukofoti",
        r"суғурта\s+мукофоти",
        r"premium",
    ],
    "declared_rate": [r"(?:страховой\s+)?тариф", r"(?:страховая\s+)?ставка", r"tarif", r"тариф", r"rate"],
    "term_days": [r"срок\s*\(?дн(?:ей|и)?\)?", r"muddat\s*\(?kun\)?", r"term\s*\(?days\)?"],
}


def check_archive(data):
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        members = z.infolist()
        if len(members) > 3000 or sum(x.file_size for x in members) > 50 * 1024 * 1024:
            raise ValueError("Архив документа слишком большой после распаковки")
        if any(x.filename.endswith("vbaProject.bin") for x in members):
            raise ValueError("Документы с макросами не принимаются")


def extract_text(data, filename, max_pages=50):
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        if not data.startswith(b"%PDF"):
            raise ValueError("Файл не является PDF")
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ValueError("PDF защищён паролем")
        if len(reader.pages) > max_pages:
            raise ValueError(f"Допустимо не более {max_pages} страниц")
        return "\n".join(p.extract_text() or "" for p in reader.pages), len(reader.pages)
    if suffix in {".docx", ".xlsx"}:
        check_archive(data)
        if suffix == ".docx":
            doc = WordDocument(io.BytesIO(data))
            text = "\n".join(
                [p.text for p in doc.paragraphs]
                + [" | ".join(c.text for c in r.cells) for t in doc.tables for r in t.rows]
            )
            if len(text) > max_pages * 5000:
                raise ValueError("Текст документа превышает лимит")
            return text, None
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        text = []
        for ws in wb:
            if ws.max_row > 10000 or ws.max_column > 100:
                raise ValueError("Excel: максимум 10 000 строк и 100 столбцов")
            for row in ws.iter_rows(values_only=True):
                text.append(" | ".join(str(v) if v is not None else "" for v in row))
                if len(text) > 10000:
                    raise ValueError("Слишком много строк Excel")
        wb.close()
        return "\n".join(text), None
    if suffix in {".png", ".jpg", ".jpeg", ".webp"}:
        with Image.open(io.BytesIO(data)) as im:
            if im.width * im.height > 25_000_000:
                raise ValueError("Фото больше 25 мегапикселей")
            im.verify()
        return "", 1
    if suffix in {".txt", ".csv"}:
        return data.decode("utf-8-sig"), None
    raise ValueError(
        "Поддерживаются PDF, DOCX, XLSX, TXT, CSV, JPG, PNG, WEBP. Старые Word/Excel сохраните как DOCX/XLSX."
    )


def parse_document(data, filename, max_pages=50):
    text, pages = extract_text(data, filename, max_pages)
    if not text.strip():
        return {
            "kind": "photo_or_scan",
            "mode": "manual",
            "pages": pages,
            "fields": {},
            "notice": "ИИ отключён. Фото/скан сохранён: изучите файл и заполните данные вручную.",
        }
    lower = text.lower()
    kind = (
        "contract"
        if re.search(r"договор|shartnoma|шартнома|contract", lower)
        else "branch_request"
        if re.search(r"запрос|филиал|so['‘’]?rov|filial", lower)
        else "report"
        if re.search(r"отч[её]т|hisobot|report", lower)
        else "document"
    )
    fields = {}
    for field, patterns in FIELDS.items():
        match = re.search("(?:" + "|".join(patterns) + r")[ \t]*[:=|\-]?[ \t]*([^\n|]*)", text, re.I)
        if not match:
            fields[field] = {"value": None, "status": "unavailable", "source": filename}
            continue
        raw = match.group(1).strip()
        value = document_number(raw, field)
        fields[field] = {
            "value": value,
            "status": "extracted"
            if value
            else "blank"
            if not raw or re.fullmatch(r"[_ .-]+", raw)
            else "needs_review",
            "source": filename,
            "excerpt": match.group(0)[:250],
        }
    text_fields = {
        "object_description": r"объект\s+страхования|sug['‘’]?urta\s+obyekti|insured\s+object",
        "insured_organization": r"страхователь|sug['‘’]?urta\s+qildiruvchi|policyholder",
        "insurer_organization": r"страховщик|sug['‘’]?urtalovchi|insurer",
        "contract_start": r"дата\s+начала|начало\s+страхования|start\s+date|boshlanish\s+sanasi",
        "contract_end": r"дата\s+окончания|окончание\s+страхования|end\s+date|tugash\s+sanasi",
    }
    for field, pattern in text_fields.items():
        match = re.search(r"(?:" + pattern + r")[ \t]*[:=|][ \t]*([^\n|]*)", text, re.I)
        raw = match.group(1).strip() if match else ""
        value = raw if raw and not re.fullmatch(r"[_ .-]+", raw) else None
        if (
            value
            and field.endswith("organization")
            and not re.search(r"\b(?:ООО|АО|ОАО|ЗАО|МЧЖ|АЖ|MChJ|AJ|LLC|JSC)\b", value, re.I)
        ):
            value = None
        if value and field.startswith("contract_"):
            value = None
            for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
                try:
                    value = datetime.strptime(raw[:10], fmt).date().isoformat()
                    break
                except ValueError:
                    pass
        fields[field] = {
            "value": value,
            "source": filename,
            "status": "extracted"
            if value
            else "unavailable"
            if not match
            else "blank"
            if not raw or re.fullmatch(r"[_ .-]+", raw)
            else "needs_review",
        }
        if match and (not field.endswith("organization") or value):
            fields[field]["excerpt"] = match.group(0)[:250]
    derive_document_term(fields, filename)
    return {
        "kind": kind,
        "mode": "rules",
        "pages": pages,
        "fields": fields,
        "notice": "Распознано правилами. Проверьте значения. Стороны-физлица автоматически не извлекаются.",
    }


def read_table(data, filename):
    try:
        return _read_table(data, filename)
    except (zipfile.BadZipFile, StopIteration, UnicodeDecodeError, KeyError, IndexError):
        raise ValueError("Файл повреждён или пуст. Используйте CSV в UTF-8 или XLSX с заголовками") from None


def _read_table(data, filename):
    if filename.lower().endswith(".csv"):
        text = data.decode("utf-8-sig")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(io.StringIO(text), dialect=dialect)
        if not reader.fieldnames or len(set(reader.fieldnames)) != len(reader.fieldnames):
            raise ValueError("Заголовки столбцов должны быть уникальны и заполнены")
        rows = list(reader)
    elif filename.lower().endswith(".xlsx"):
        check_archive(data)
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.active
        if ws.max_row > 10001 or ws.max_column > 100:
            raise ValueError("Максимум 10 000 строк и 100 столбцов")
        values = iter(ws.values)
        headers = [str(x).strip() if x is not None else "" for x in next(values)]
        if len(set(headers)) != len(headers):
            raise ValueError("Заголовки столбцов должны быть уникальны")
        rows = [dict(zip(headers, row, strict=False)) for row in values if any(v is not None for v in row)]
        wb.close()
    else:
        raise ValueError("Для импорта используйте CSV или XLSX")
    if not rows or len(rows) > 10000:
        raise ValueError("Нужно от 1 до 10 000 строк")
    return [{str(k).strip(): v for k, v in row.items() if v not in (None, "")} for row in rows]
