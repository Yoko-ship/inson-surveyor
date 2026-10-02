import io
import os
import re
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from surveyor.config import settings


def ai_plain(text):
    """Format model prose for exports; never apply this to source quotations."""
    for marker in ("**", "__", "`", "*"):
        pattern = re.escape(marker) + r"([^\n]+?)" + re.escape(marker)
        text = re.sub(pattern, r"\1", text)
    return re.sub(r"(?m)^#{1,6}\s+", "", text)


LABELS = {
    "ru": [
        "Сюрвейерский акт",
        "1. Объект и документы",
        "2. Оценка стоимости",
        "3. Тариф и страховая премия",
        "4. Аналитика риска",
        "5. Заключение и оговорки",
        "Данные недоступны",
        "Подлежит подтверждению андеррайтером · Не является кредитным скорингом",
    ],
    "uz": [
        "Syurveyer dalolatnomasi",
        "1. Obyekt va hujjatlar",
        "2. Qiymatni baholash",
        "3. Tarif va sug‘urta mukofoti",
        "4. Xatar tahlili",
        "5. Xulosa va shartlar",
        "Ma’lumot mavjud emas",
        "Anderrayter tasdig‘i talab qilinadi · Kredit skoringi emas",
    ],
    "en": [
        "Survey report",
        "1. Object and documents",
        "2. Property valuation",
        "3. Rate and insurance premium",
        "4. Risk analysis",
        "5. Conclusion and clauses",
        "Data unavailable",
        "Subject to underwriter confirmation · Not a credit score",
    ],
}


def sections(report):
    from surveyor.regions import REGIONS, region_code
    from surveyor.report_text import label, warning

    s = report.snapshot
    lang = s["language"]
    labels = LABELS[lang]

    def tr(key):
        return label(key, lang)

    missing = labels[6]

    def value(v):
        if v is None or v == "":
            return missing
        if isinstance(v, bool):
            return {"ru": ("нет", "да"), "uz": ("yo‘q", "ha"), "en": ("no", "yes")}[lang][v]
        return str(v)

    def pairs(data, keys):
        return [
            f"{tr(k)}: {tr(str(data[k])) if k in {'method', 'status', 'rate_type', 'basis', 'risk_level', 'comparison', 'evidence_kind'} and data.get(k) else value(data.get(k))}"
            for k in keys
        ]

    def comparable_lines(rows):
        lines = []
        for c in rows:
            lines.append(c["label"])
            lines += pairs(
                c,
                [
                    "price",
                    "original_price",
                    "price_uzs",
                    "date",
                    "source",
                    "evidence_kind",
                    "transaction_reference",
                    "edit_reason",
                ],
            )
            if c.get("reason"):
                lines.append(tr("reason") + ": " + warning(c["reason"], lang))
            if c.get("exchange_source"):
                fx = c["exchange_source"]
                lines.append(
                    f"{c['currency']} → UZS: {fx['value']} · {fx['source_url']} · {fx['observation_date']}"
                )
        return lines or [tr("none")]

    object_lines = [
        s["title"],
        f"ID: {report.id}",
        f"{tr('date')}: {report.created_at.isoformat()}",
        f"{tr('author')}: {s['author']['name']}",
    ]
    inputs = dict(s["inputs"])
    code = region_code(inputs["region"])
    inputs["region"] = next(
        (r[{"ru": 1, "uz": 2, "en": 3}[lang]] for r in REGIONS if r[0] == code), inputs["region"]
    )
    object_lines += pairs(
        inputs,
        [
            "insured_sum",
            "object_value",
            "region",
            "term_days",
            "contract_start",
            "contract_end",
            "object_description",
        ],
    )
    if inputs.get("contract_start") and inputs.get("contract_end"):
        object_lines.append(tr("term_date_convention"))
    object_lines.append(tr("manual_input"))
    for doc in s["documents"]:
        object_lines.append(f"{tr('document')}: {doc['filename']} · SHA256 {doc['sha256']}")
        for k, item in doc["extracted"]["fields"].items():
            object_lines.append(
                f"{tr(k)}: {value(item.get('value'))} · {tr(item['status'])} · {item['source']}"
            )
            if item.get("excerpt"):
                object_lines.append(item["excerpt"])
            if item.get("original") is not None:
                object_lines.append(
                    f"{tr('corrections')}: {value(item['original'])} → {value(item.get('value'))}"
                )
        for review in doc["extracted"].get("ai_reviews", []):
            object_lines.append(
                f"{tr('ai_review')}: {review['reviewer_name']} ({review['reviewed_by']}) · {review['reviewed_at']}"
            )
            object_lines.append(
                f"{review['provider']} · {review['model']} · {tr('ai_config_revision')}: {review['config_revision']}"
            )
            object_lines.append(f"{tr('reason')}: {review['reason']}")
            for row in review["fields"]:
                object_lines.append(
                    f"{tr(row['field'])} · {tr('ai_' + row['decision'])}: "
                    f"{value(row['value'])} → {value(row['reviewed_value'])}"
                )
                if row.get("quote"):
                    object_lines.append(f"{tr('ai_quote')}: {row['quote']}")
                    if not row.get("quote_found_in_text"):
                        object_lines.append(tr("ai_quote_visual"))
    assistance = s.get("assistance", {})
    if assistance.get("answers") or assistance.get("reviews"):
        object_lines.append(tr("inspection_guidance"))
    for key, answer in assistance.get("answers", {}).items():
        object_lines.append(f"{assistance.get('question_labels', {}).get(key, tr(key))}: {answer}")
    for review in assistance.get("reviews", []):
        object_lines.append(
            f"{tr('ai_review')}: {review['reviewer_name']} · {review['reviewed_at']} · {review['provider']} · {review['config_revision']}"
        )
        object_lines.append(f"{tr('reason')}: {review['reason']}")
        if review.get("context", {}).get("product"):
            object_lines.append(f"{tr('product')}: {review['context']['product']['version_id']}")
        if review["stale"]:
            object_lines.append(tr("ai_review_stale"))
        for row in review["fields"]:
            object_lines.append(f"{tr('ai_' + row['decision'])}: {ai_plain(row['text'])}")
            if row.get("reviewed_text"):
                object_lines.append(f"{tr('corrections')}: {ai_plain(row['reviewed_text'])}")
            for cite in row["citations"]:
                object_lines.append(f"{cite.get('filename') or cite['source_id']}: {cite['quote']}")
                if not cite["quote_found_in_text"]:
                    object_lines.append(tr("ai_quote_visual"))
    object_lines.append(tr("corrections"))
    object_lines += [f"{tr(k)}: {v}" for k, v in inputs.get("overrides", {}).items()] or [tr("none")]
    object_lines.append(tr("reason") + ": " + value(inputs.get("override_reason")))
    object_lines.append(tr("conflicts"))
    for conflict in s["conflicts"]:
        object_lines.append(
            f"{tr(conflict['field'])}: "
            + " / ".join(f"{x['source']} = {x['value']}" for x in conflict["sources"])
            + f"; {tr('manual_input')}: {value(conflict['entered'])}"
        )
    if not s["conflicts"]:
        object_lines.append(tr("none"))
    borrower = inputs.get("borrower")
    if borrower:
        object_lines.append(tr("borrower"))
        object_lines += pairs(
            borrower, ["organization_name", "bureau_name", "report_date", "score", "score_scale", "summary"]
        )
        evidence = next((d for d in s["documents"] if d["id"] == borrower["document_id"]), None)
        if evidence:
            object_lines.append(f"{tr('source')}: {evidence['filename']} · SHA256 {evidence['sha256']}")
        object_lines.append(tr("external_credit_score"))
    v = s["valuation"]
    value_lines = pairs(
        v,
        [
            "method",
            "estimate",
            "minimum",
            "maximum",
            "raw_median",
            "deviation_percent",
            "second_method_deviation_percent",
            "status",
        ],
    )
    value_lines += comparable_lines(v["comparables"]) + [tr("excluded")] + comparable_lines(v["rejected"])
    if v.get("evidence_missing"):
        value_lines.append(tr("missing_evidence") + ": " + ", ".join(tr(k) for k in v["evidence_missing"]))
    policy = v.get("policy", {})
    value_lines += [tr("policy")] + pairs(
        policy, ["outlier_low", "outlier_high", "tolerance", "requires_appraiser", "approved_by"]
    )
    if not policy.get("approved_by"):
        value_lines.append(tr("unapproved"))
    for key in [
        "purchase_price",
        "purchase_source",
        "purchase_date",
        "depreciation_percent",
        "appraiser_value",
        "appraiser_source",
        "appraiser_date",
        "second_method_value",
        "second_method_source",
        "second_method_date",
    ]:
        if inputs.get(key) is not None:
            value_lines += pairs(inputs, [key])
    c = s["calculation"]
    rate_lines = [
        f"{tr('product')}: {s['product']['code']} — {s['product']['name']}",
        f"{tr('policy_version')}: {s['product']['version_id']}; {tr('effective_from')}: {s['product']['effective_from']}",
    ]
    rate_lines += pairs(
        c,
        [
            "rate_type",
            "minimum_rate",
            "recommended_rate",
            "annualized_rate",
            "annual_market_rate",
            "premium",
            "formula",
            "premium_discrepancy",
            "comparison",
        ],
    )
    if s["product"].get("normative_source"):
        rate_lines.append(s["product"]["normative_source"])
    source = s["product"].get("tariff_policy")
    if source:
        rate_lines.append(
            f"{tr('tariff_source')}: {source['source_filename']} · {source['page']} · SHA256 {source['sha256']}"
        )
        rate_lines.append(f"{source['code']} · {source.get('variant') or '—'}")
        rate_lines += pairs(
            s["product"],
            [
                "policy_basis_reference",
                "policy_terms_reference",
                "policy_approval_reference",
                "agent_commission_percent",
            ],
        )
    if s.get("rnp_classification"):
        rnp = s["rnp_classification"]
        rate_lines.append(f"{tr('rnp_group')}: {rnp['group'] or missing}")
        rate_lines += [rnp["reason"], rnp["source_url"], rnp["rule_version"], tr("rnp_scope")]
    if c.get("reason"):
        rate_lines.append(warning(c["reason"], lang))
    factors = c.get("factor_pricing")
    if factors:
        rate_lines += [
            tr("factor_pricing"),
            f"{tr('template')}: {factors['template_id']}",
            f"{tr('source')}: {factors['source_id']} · SHA256 {factors['source_sha256']}",
            f"{tr('factor_status')}: {tr('factor_' + factors['status'])}",
            f"{tr('approved_by')}: {value(factors['approved_by'])} · {value(factors['approved_at'])}",
            f"{tr('factor_multiplier')}: {factors['multiplier']}",
            f"{tr('factor_proposal')}: {factors['proposed_multiplier']}",
            tr("factor_formula"),
            factors["rationale"],
        ]
        for row in factors["rows"]:
            if row["choice"] == "unanswered":
                continue
            rate_lines.append(
                f"{row['label']} · {tr('factor_' + row['choice'])} · "
                f"{tr('factor_coefficient')}: {value(row['coefficient'])} · "
                f"{tr('factor_applied')}: {row['applied_coefficient']} · {tr('factor_page')}: {row['page']}"
            )
            if row["source_condition"]:
                rate_lines.append(row["source_condition"])
            rate_lines.append(row["evidence"])
        if factors.get("calibration"):
            cal = factors["calibration"]
            rate_lines += [
                f"{tr('calibration')}: {cal['method']} · {cal['data_id']} · SHA256 {cal['data_sha256']}",
                tr("factor_method_note"),
                cal["rationale"],
            ]
            for estimate in cal["estimates"]:
                rate_lines.append(
                    f"{estimate['factor_id']} · {tr('factor_' + estimate['choice'])}: "
                    f"{estimate['raw']} → {estimate['coefficient']} · {estimate['years']}"
                )
        else:
            rate_lines.append(tr("factor_uncalibrated"))
    risk_lines = pairs(
        c, ["risk_score", "risk_level", "risk_multiplier", "regional_adjustment", "loss_adjustment"]
    )
    risk_lines += [warning(w, lang) for w in c.get("warnings", [])]
    risk_lines.append(tr("losses"))
    for loss in s["losses"]:
        risk_lines += pairs(
            loss, ["year", "claims", "payments", "premiums", "contracts", "loss_ratio", "frequency"]
        )
    if not s["losses"]:
        risk_lines.append(missing)
    if s.get("calibration"):
        risk_lines.append(tr("calibration") + ": " + s["calibration"]["id"])
        risk_lines += pairs(s["calibration"], ["loss_ratio", "approved_by"])
        risk_lines.append(s["calibration"]["rationale"])
    if s.get("template"):
        risk_lines.append(tr("template") + ": " + s["template"]["id"])
        risk_lines.append(
            tr("risk_shares")
            + ": "
            + "; ".join(f"{k}: {v}%" for k, v in s["template"].get("risk_shares", {}).items())
        )
    # Show relevant regional/class data and only the FX currencies actually used in valuation.
    currencies = {x.get("currency", "UZS") for x in inputs.get("comparables", [])}
    for indicator in s["indicators"]:
        if indicator["metric"].startswith("fx_") and indicator["metric"][3:] not in currencies:
            continue
        if indicator.get("reference_only"):
            indicator_region = region_code(indicator.get("region", "all"))
            indicator_region = next(
                (r[{"ru": 1, "uz": 2, "en": 3}[lang]] for r in REGIONS if r[0] == indicator_region),
                indicator_region,
            )
            risk_lines.append(
                f"{tr('region')}: {indicator_region} · {tr('class_code')}: {indicator.get('class_code', 'all')}"
            )
            risk_lines.append(tr("reference_only"))
        if indicator.get("subject", "market") != "market":
            risk_lines.append(indicator["subject"])
        if indicator.get("note"):
            risk_lines.append(indicator["note"])
        risk_lines.append(
            f"{tr(indicator['metric'])}: {indicator['value']} {indicator['unit']} · {indicator['period']} · {indicator['source_url']} · {indicator['fetched_at']}"
            + (" · " + tr("stale") if indicator["stale"] else "")
        )
        if indicator.get("quote"):
            risk_lines += pairs(indicator["quote"], ["rate", "basis", "term_days", "coverage"])
    for reference in s.get("references", []):
        risk_lines += [
            reference["title"],
            reference["source_url"],
            value(reference.get("observation_date")),
            reference["text"],
            "SHA256 " + reference["sha256"],
        ]
        if reference["stale"]:
            risk_lines.append(tr("stale"))
    conclusion = [
        tr(k)
        for k in [
            "underwriter",
            "not_credit",
            "human_review",
            "check_values",
            "check_conflicts",
            "check_value",
            "check_rates",
            "check_approval",
        ]
    ]
    if s["clauses"]:
        conclusion += [tr("original_language")] + s["clauses"]
    if factors and factors["clarify"]:
        conclusion.append(tr("factor_clarify"))
        conclusion += [f"{row['label']}: {tr('factor_' + row['reason'])}" for row in factors["clarify"]]
    return labels, [
        (labels[1], object_lines),
        (labels[2], value_lines),
        (labels[3], rate_lines),
        (labels[4], risk_lines),
        (labels[5], conclusion),
    ]


def export_docx(report):
    labels, blocks = sections(report)
    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10)
    doc.add_heading(labels[0], 0)
    for title, lines in blocks:
        doc.add_heading(title, 1)
        for line in lines:
            doc.add_paragraph(line)
    doc.sections[0].footer.paragraphs[0].text = labels[7]
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def export_pdf(report):
    labels, blocks = sections(report)
    candidates = [
        settings.pdf_font_path,
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        str(Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / "arial.ttf"),
    ]
    font_path = next((p for p in candidates if p and Path(p).is_file()), None)
    if not font_path:
        raise ValueError("Установите шрифт DejaVu Sans или задайте PDF_FONT_PATH")
    if "Surveyor" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("Surveyor", font_path))
    body = ParagraphStyle("BodyRu", fontName="Surveyor", fontSize=9, leading=13, spaceAfter=6, wordWrap="CJK")
    heading = ParagraphStyle(
        "HeadRu",
        parent=body,
        fontSize=14,
        leading=19,
        textColor=colors.HexColor("#17463e"),
        spaceBefore=15,
        spaceAfter=10,
    )
    buffer = io.BytesIO()
    pdf = SimpleDocTemplate(
        buffer, topMargin=18 * mm, bottomMargin=22 * mm, rightMargin=18 * mm, leftMargin=18 * mm
    )
    story = [Paragraph(escape(labels[0]), heading)]
    for title, lines in blocks:
        story.append(Paragraph(escape(title), heading))
        for line in lines:
            story.append(Paragraph(escape(line).replace("\n", "<br/>"), body))
        story.append(Spacer(1, 5 * mm))

    def footer(canvas, doc):
        canvas.setFont("Surveyor", 7)
        canvas.drawString(18 * mm, 12 * mm, labels[7])
        canvas.drawRightString(190 * mm, 8 * mm, str(doc.page))

    pdf.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()
