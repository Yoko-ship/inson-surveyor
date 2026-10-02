"""Bounded NAPP reference tables; never convert portfolio aggregates to quote rates."""

import io
import re
from datetime import date, timedelta
from decimal import Decimal
from urllib.parse import unquote

from openpyxl import load_workbook

from surveyor.napp import parse_napp
from surveyor.regions import REGIONS
from surveyor.schemas import IndicatorInput

REGIONAL_NOTE = (
    "Только справочно: учёт по головным офисам и подразделениям может искажать регион объекта. "
    "Выплаты / премии — агрегат за период, не актуарная убыточность и не коэффициент тарифа."
)
PORTFOLIO_NOTE = "Статистика портфеля за период; не годовой тариф и не продуктовая история убытков."
REGION_NAMES = [
    "Qoraqalpog‘iston Respublikasi",
    "Andijon viloyati",
    "Buxoro viloyati",
    "Jizzax viloyati",
    "Qashqadaryo viloyati",
    "Navoiy viloyati",
    "Namangan viloyati",
    "Samarqand viloyati",
    "Surxandaryo viloyati",
    "Sirdaryo viloyati",
    "Toshkent viloyati",
    "Farg`ona viloyati",
    "Xorazm viloyati",
    "Toshkent sh.",
]


def normalized(value):
    return " ".join(str(value or "").split()).replace("`", "‘").replace("'", "‘")


def amount(value):
    if value in (None, "", "-", "—"):
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
        raise ValueError("NAPP reference numeric format changed")
    result = Decimal(str(value))
    if not result.is_finite() or result < 0:
        raise ValueError("NAPP reference invalid amount")
    return result


def parse_napp_reference(content, url):
    # Reuse archive, unit, date and 18-class validation before opening additional sheets.
    parse_napp(content, url)
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:

        def rows(name):
            return list(wb[name].iter_rows(max_row=100, max_col=34, values_only=True))

        class_rows = rows("1.4")
        dates = class_rows[3][1:3]
        end = dates[1].date() - timedelta(days=1)
        path_period = re.search(r"/opendata/(20\d{2})/([1-4])Q/", unquote(url))
        if not path_period or (int(path_period[1]), int(path_period[2])) != (
            end.year,
            (end.month - 1) // 3 + 1,
        ):
            raise ValueError("NAPP URL and workbook periods differ")
        result = []

        def check_periods(table, row, col):
            if tuple(table[row][col : col + 2]) != tuple(dates):
                raise ValueError("NAPP reference periods changed")

        def check_header(table, row, col, expected):
            if normalized(table[row][col]) != normalized(expected):
                raise ValueError(f"NAPP reference header changed: {expected}")

        def check_units(table):
            if not any("mln." in str(v) and "so‘m" in normalized(v) for row in table[:5] for v in row):
                raise ValueError("NAPP reference monetary unit changed")

        def add(metric, value, subject="market", region="all", cls="all", unit="UZS", offset=1, note=None):
            value = amount(value)
            if value is None:
                return
            observed = dates[offset].date() - timedelta(days=1)
            if observed > date.today():
                raise ValueError("NAPP future period")
            result.append(
                IndicatorInput(
                    metric="napp_ref_" + metric,
                    value=value * (1000000 if unit == "UZS" else 1),
                    subject=subject,
                    region=region,
                    class_code=cls,
                    unit=unit,
                    period=f"{observed.year}-01-01/{observed.isoformat()}",
                    observation_date=observed,
                    source_url=url,
                    stale_days=200,
                    reference_only=True,
                    note=note or (REGIONAL_NOTE if region != "all" else PORTFOLIO_NOTE),
                )
            )

        # Bundles retain their combined class scope; never distribute them to individual classes.
        bundles = 0
        for row in class_rows:
            match = re.fullmatch(r"(\d+(?:,\d+)+) klasslar", str(row[0]).strip())
            if not match:
                continue
            numbers = [int(v) for v in match[1].split(",")]
            if len(set(numbers)) != len(numbers) or not all(1 <= n <= 18 for n in numbers):
                raise ValueError("NAPP bundle class list changed")
            cls = "bundle_" + "_".join(map(str, sorted(numbers)))
            for offset in (0, 1):
                for metric, col in [
                    ("bundle_premiums", 1),
                    ("bundle_payments", 4),
                    ("bundle_liabilities", 7),
                ]:
                    add(
                        metric,
                        row[col + offset],
                        cls=cls,
                        offset=offset,
                        note="Комплексное страхование: единая сумма по набору классов, без распределения между ними.",
                    )
            bundles += 1
        if not bundles:
            raise ValueError("NAPP bundle table missing")

        # Company totals and claims. Exact company names remain separate identities.
        for sheet, header_row, date_row, start, col, header, metric in [
            ("2.1", 3, 4, 6, 11, "Umumiy sug`urta mukofotlari", "company_premiums"),
            ("2.5", 3, 4, 6, 2, "Sug‘urta to‘lovlari", "company_payments"),
        ]:
            table = rows(sheet)
            check_header(table, header_row, col, header)
            check_periods(table, date_row, col)
            check_units(table)
            count = 0
            for row in table[start:]:
                if not isinstance(row[0], (int, float)) or not isinstance(row[1], str):
                    continue
                for offset in (0, 1):
                    add(metric, row[col + offset], subject=row[1].strip(), offset=offset)
                count += 1
            if not count:
                raise ValueError("NAPP company table missing")

        region_map = {normalized(name): REGIONS[i + 1][0] for i, name in enumerate(REGION_NAMES)}
        region_map.update(
            {
                normalized(name.removesuffix(" viloyati")): REGIONS[i + 1][0]
                for i, name in enumerate(REGION_NAMES[:-1])
            }
        )
        region_map["Toshkent shahri"] = "1726"
        for sheet, date_row, label, start in [
            ("2.10", 3, "Sug‘urta kompaniyalari nomi", 7),
            ("3.5", 2, "Hudud", 6),
        ]:
            table = rows(sheet)
            check_header(table, date_row, 1, label)
            if (table[date_row][2], table[date_row][9]) != tuple(dates):
                raise ValueError("NAPP claims periods changed")
            for offset, base in [(0, 2), (1, 9)]:
                check_header(table, date_row + 1, base, "Kelib tushgan sug‘urta da'volarning umumiy soni")
                check_header(table, date_row + 2, base + 1, "To‘langan")
                check_header(table, date_row + 2, base + 3, "Rad etilgan")
                check_header(table, date_row + 1, base + 5, "Qoniqtirilmagan sug‘urta da'volari soni")
                for row in table[start:]:
                    if not isinstance(row[0], (int, float)) or not isinstance(row[1], str):
                        continue
                    region = region_map.get(normalized(row[1])) if sheet == "3.5" else "all"
                    if region is None:
                        raise ValueError("NAPP unknown claims region")
                    subject = row[1].strip() if sheet == "2.10" else "market"
                    for metric, delta in [
                        ("claims_received", 0),
                        ("claims_paid", 1),
                        ("claims_refused", 3),
                        ("claims_unsettled", 5),
                    ]:
                        value = amount(row[base + delta])
                        if value is not None and value != value.to_integral_value():
                            raise ValueError("NAPP fractional claim count")
                        add(metric, value, subject=subject, region=region, unit="count", offset=offset)

        for sheet, date_row, start, metric in [
            ("3.1", 5, 7, "regional_premiums"),
            ("3.2", 4, 6, "regional_payments"),
        ]:
            table = rows(sheet)
            check_periods(table, date_row, 2)
            check_units(table)
            check_header(table, date_row - 2, 1, "Hudud")
            check_header(
                table,
                date_row - 2,
                2,
                "Umumiy sug‘urta mukofotlari" if sheet == "3.1" else "Sug‘urta to‘lovlari",
            )
            for row in table[start:]:
                if not isinstance(row[0], (int, float)):
                    continue
                region = region_map.get(normalized(row[1]))
                if region is None:
                    raise ValueError("NAPP unknown region")
                for offset in (0, 1):
                    add(metric, row[2 + offset], region=region, offset=offset)

        # Current-period company subdivisions; period reconciles with company totals below.
        subdivisions = {}
        for sheet, metric, header in [
            ("2.12", "premiums", "Sug`urta mukofotlari"),
            ("2.13", "payments", "Sug‘urta to‘lovlari"),
        ]:
            table = rows(sheet)
            check_units(table)
            for i, expected in enumerate(REGION_NAMES):
                check_header(table, 3, 4 + 2 * i, expected)
                check_header(table, 4, 4 + 2 * i, header)
            matches = [r for r in table[6:] if re.fullmatch(r'"INSON"\s+AJ', str(r[1]).strip())]
            if len(matches) != 1:
                raise ValueError("NAPP INSON subdivision row missing or duplicated")
            row = matches[0]
            total = amount(row[2])
            market_total = amount(table[5][2])
            company = next(
                (
                    x
                    for x in result
                    if x.subject == '"INSON" AJ'
                    and x.metric == f"napp_ref_company_{metric}"
                    and x.observation_date == end
                ),
                None,
            )
            if total is None or company is None or abs(total * 1000000 - company.value) > Decimal("1"):
                raise ValueError("NAPP company and subdivision totals differ")
            if market_total:
                add(
                    f"inson_{metric}_market_share",
                    total / market_total,
                    subject='"INSON" AJ',
                    unit="ratio",
                    note="Доля INSON в общем рынке, включая страхование жизни; период совпадает с отчётом.",
                )
            values = {"all": total}
            for i, region in enumerate(REGIONS[1:]):
                value = amount(row[4 + 2 * i])
                values[region[0]] = value
                add(
                    f"inson_subdivision_{metric}",
                    value,
                    subject='"INSON" AJ',
                    region=region[0],
                    note=REGIONAL_NOTE,
                )
            subdivisions[metric] = values
        total_p = subdivisions["premiums"]["all"]
        total_l = subdivisions["payments"]["all"]
        national_ratio = total_l / total_p if total_p else None
        for region, premiums in subdivisions["premiums"].items():
            payments = subdivisions["payments"].get(region)
            if premiums and payments is not None:
                ratio = payments / premiums
                add(
                    "inson_payment_premium_ratio",
                    ratio,
                    subject='"INSON" AJ',
                    region=region,
                    unit="ratio",
                    note=REGIONAL_NOTE,
                )
                if region != "all" and national_ratio:
                    add(
                        "inson_region_to_company_ratio",
                        ratio / national_ratio,
                        subject='"INSON" AJ',
                        region=region,
                        unit="ratio",
                        note=REGIONAL_NOTE,
                    )
        return result, ["napp_reference_v1", "company", "region", "bundle", "reference_only"]
    finally:
        wb.close()
