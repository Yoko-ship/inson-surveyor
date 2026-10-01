"""NAPP official class statistics. Aggregate premiums/liabilities are NOT annual tariff quotes."""

import calendar
import io
import re
import zipfile
from datetime import date, datetime, timedelta
from decimal import Decimal
from html.parser import HTMLParser
from urllib.parse import unquote, urljoin

from openpyxl import load_workbook

from surveyor.schemas import IndicatorInput

NAPP_INDEX = "https://napp.uz/en/pages/statistics-and-analysis-for-im"


class DownloadLinks(HTMLParser):
    def __init__(self):
        super().__init__()
        self.links = []

    def handle_starttag(self, tag, attrs):
        if tag == "a":
            href = dict(attrs).get("href", "")
            if href.lower().endswith(".xlsx"):
                self.links.append(urljoin(NAPP_INDEX, href))


def latest_download(content):
    parser = DownloadLinks()
    parser.feed(content.decode("utf-8"))
    choices = []
    for url in parser.links:
        match = re.search(r"/opendata/(20\d{2})/([1-4])Q/", unquote(url))
        if not match:
            continue
        year, quarter = map(int, match.groups())
        month = quarter * 3
        end = date(year, month, calendar.monthrange(year, month)[1])
        if end <= date.today():
            choices.append((end, url))
    if not choices:
        raise ValueError("NAPP download index schema changed")
    return max(choices)[1]


def parse_napp(content, url):
    # The official workbook includes formatting to row 1,048,576 on unused sheets.
    # Read ONLY the bounded class sheet; do not iterate those sheets or relax upload limits.
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        members = archive.infolist()
        if len(members) > 1000 or sum(m.file_size for m in members) > 200 * 1024 * 1024:
            raise ValueError("NAPP workbook exceeds archive limits")
        if any(
            m.filename.endswith("vbaProject.bin")
            or (m.filename.endswith(("sharedStrings.xml", "styles.xml")) and m.file_size > 5 * 1024 * 1024)
            for m in members
        ):
            raise ValueError("Unexpected NAPP workbook content")
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        ws = wb["1.4"]
        rows = list(ws.iter_rows(min_row=1, max_row=40, max_col=10, values_only=True))
        if (
            rows[2][1] != "Umumiy sug‘urta mukofotlari"
            or rows[2][4] != "Sug‘urta to‘lovlari"
            or rows[2][7] != "Sug‘urta majburiyatlari"
        ):
            raise ValueError("NAPP class headers changed")
        if "mln." not in str(rows[4][1]) or "mln." not in str(rows[4][4]):
            raise ValueError("NAPP amount units changed")
        dates = [rows[3][1], rows[3][2]]
        if not all(isinstance(d, datetime) and d.day == 1 and d.month in {1, 4, 7, 10} for d in dates):
            raise ValueError("NAPP reporting period changed")
        if dates[1] <= dates[0] or rows[3][4:6] != tuple(dates) or rows[3][7:9] != tuple(dates):
            raise ValueError("NAPP periods do not reconcile")
        result = []
        classes = set()
        for row in rows[5:]:
            match = re.match(r"^(\d+)-klass\s", str(row[0]))
            if not match:
                continue
            number = int(match[1])
            classes.add(number)
            cls = {3: "vehicle", 8: "fire", 9: "property"}.get(number, f"class_{number}")
            for offset, reporting_date in enumerate(dates):
                end = reporting_date.date() - timedelta(days=1)
                if end > date.today():
                    raise ValueError("NAPP future reporting period")
                base = {
                    "class_code": cls,
                    "period": f"{end.year}-01-01/{end.isoformat()}",
                    "source_url": url,
                    "observation_date": end,
                    "stale_days": 200,
                }
                amounts = []
                for metric, col in [
                    ("market_premiums", 1),
                    ("market_payments", 4),
                    ("market_liabilities", 7),
                ]:
                    cell = row[col + offset]
                    if not isinstance(cell, (int, float)) or cell < 0:
                        raise ValueError("NAPP numeric cell changed")
                    amount = Decimal(str(cell)) * 1000000
                    amounts.append(amount)
                    result.append(IndicatorInput(**base, metric=metric, value=amount, unit="UZS"))
                if amounts[0] > 0:
                    result.append(
                        IndicatorInput(
                            **base, metric="market_loss_ratio", value=amounts[1] / amounts[0], unit="ratio"
                        )
                    )
        if classes != set(range(1, 19)):
            raise ValueError("NAPP class list changed")
        return result, [
            "napp_class_1_18",
            "premiums_million_UZS",
            "payments_million_UZS",
            "liabilities_million_UZS",
        ]
    finally:
        wb.close()
